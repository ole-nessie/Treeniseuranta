#!/usr/bin/env python3
"""
Convert Harjoitusohjelma_exp.xlsx to per-year v2 JSON files (export/<year>.json)
for the Treeniseuranta web app. Columns are mapped by header label, so all sheet
layouts from 2013 onwards are handled.

Usage: python3 convert_xlsx.py [workbook.xlsx]
"""

import json
import os
import sys
from datetime import datetime, time, timedelta, timezone

import openpyxl

XLSX = sys.argv[1] if len(sys.argv) > 1 else 'Harjoitusohjelma_exp.xlsx'
OUT_DIR = 'export'

HEADER_MAP = {
    'Sairaspäivä': 'sick',
    'Kilpailukalenteri': 'compCal',
    'Harjoitus': '_training',
    'Nuolien määrä': 'arrows',
    'Kilpailu': 'competition',
    '1. kierros': 'round1',
    '2. kierros': 'round2',
    'Tulos': 'result',
    'Matka': 'distanceM',
    'Matka [m]': 'distanceM',
    'Muistiinpanoja': 'notes',
    'Aerobinen': 'aerobicH',
    'Aerobinen [h]': 'aerobicH',
    'Laji': 'aerobicNotes',
    'Uinti [km]': 'swimKm',
    'Juoksu [km]': 'runKm',
    'Kävely [km]': 'walkKm',
    'Pyöräily [km]': 'cycleKm',
    'Pyöräily km': 'cycleKm',
    'Hiihto [km]': 'skiKm',
    'Hiihto km': 'skiKm',
    'km': '_km',
    'Aerobinen, muistiinpanot': 'aerobicNotes',
    'Venyttely': 'careH',
    'Lihashuolto [h]': 'careH',
    'Kohde': 'careNotes',
    'Lihashuolto, muistiinpanot': 'careNotes',
    'Sali': 'gymH',
    'Lihaskunto [h]': 'gymH',
    'Lihaskunto, muistiinpanot': 'gymNotes',
    'Kiipeily [h]': 'climbH',
    'Kiipeily, muistiinpanot': 'climbNotes',
}

HOURS = {'aerobicH', 'careH', 'gymH', 'climbH'}
INTS = {'arrows', 'round1', 'round2', 'result', 'distanceM'}
NUMS = {'swimKm', 'runKm', 'walkKm', 'cycleKm', 'skiKm'}
SUMMED = ['sick', 'arrows', 'aerobicH', 'swimKm', 'runKm', 'walkKm', 'cycleKm', 'skiKm',
          'careH', 'gymH', 'climbH']

SPORT_KM = {'uinti': 'swimKm', 'juoksu': 'runKm', 'kävely': 'walkKm',
            'hiihto': 'skiKm', 'pyöräily': 'cycleKm'}


def fmt_hours(val):
    if isinstance(val, timedelta):
        minutes = round(val.total_seconds() / 60)
    elif isinstance(val, time):
        minutes = val.hour * 60 + val.minute + round(val.second / 60)
    elif isinstance(val, datetime):
        minutes = val.hour * 60 + val.minute
    elif isinstance(val, (int, float)):
        minutes = round(val * 24 * 60)  # Excel day fraction
    else:
        return None
    if minutes <= 0:
        return None
    return f'{minutes // 60}:{minutes % 60:02d}'


def hours_to_min(s):
    h, m = s.split(':')
    return int(h) * 60 + int(m)


def is_empty(v):
    return v is None or v == '' or v == 0 or (isinstance(v, str) and not v.strip())


def find_header(ws):
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=6, values_only=True), start=1):
        if row and 'Päivämäärä' in row:
            return idx, list(row)
    raise ValueError(f'No header row found in {ws.title}')


def build_colmap(header):
    colmap = {}
    date_col = header.index('Päivämäärä')
    for i, label in enumerate(header):
        if not isinstance(label, str):
            continue
        key = HEADER_MAP.get(label.strip())
        if key and key not in colmap.values():
            colmap[i] = key
    # Old sheets: the unlabelled column right after the gym hours holds the gym notes.
    for i, key in list(colmap.items()):
        if key == 'gymH' and 'gymNotes' not in colmap.values():
            nxt = i + 1
            if nxt < len(header) and header[nxt] is None:
                colmap[nxt] = 'gymNotes'
    return date_col, colmap


def convert_value(key, val):
    if is_empty(val):
        return None
    if key == 'sick':
        return True
    if key in HOURS:
        return fmt_hours(val)
    if key in INTS:
        return int(round(val)) if isinstance(val, (int, float)) else None
    if key in NUMS or key == '_km':
        return round(float(val), 2) if isinstance(val, (int, float)) else None
    if isinstance(val, (datetime, time, timedelta)):
        return None
    return str(val).strip()


def parse_sheet(ws, year):
    header_row, header = find_header(ws)
    date_col, colmap = build_colmap(header)
    days = {}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if date_col >= len(row) or not isinstance(row[date_col], datetime):
            continue
        dt = row[date_col]
        if dt.year != year:
            continue
        day = {}
        for i, key in colmap.items():
            if i >= len(row):
                continue
            v = convert_value(key, row[i])
            if v is not None:
                day[key] = v
        training = day.pop('_training', None)
        if training:
            day['notes'] = (f'Harjoitus: {training}\n' + day.get('notes', '')).strip()
        km = day.pop('_km', None)
        if km:
            sport = (day.get('aerobicNotes') or '').strip().lower()
            day[SPORT_KM.get(sport, 'cycleKm')] = km
        if day:
            days[dt.strftime('%Y-%m-%d')] = day
    return days


def totals(days):
    t = {}
    for d in days.values():
        for k in SUMMED:
            if k not in d:
                continue
            if k == 'sick':
                t[k] = t.get(k, 0) + 1
            elif k in HOURS:
                t[k] = t.get(k, 0) + hours_to_min(d[k])
            else:
                t[k] = t.get(k, 0) + d[k]
    out = []
    for k in SUMMED:
        if k not in t:
            continue
        v = t[k]
        if k in HOURS:
            v = f'{v // 60}:{v % 60:02d}'
        elif isinstance(v, float):
            v = round(v, 2)
        out.append(f'{k}={v}')
    return ', '.join(out)


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True, read_only=True)
    os.makedirs(OUT_DIR, exist_ok=True)
    now = datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')
    for name in wb.sheetnames:
        if 'Harjoitukset' not in name:
            continue
        year = int(name.split()[0])
        days = parse_sheet(wb[name], year)
        data = {'version': 2, 'year': year, 'updated': now, 'days': days}
        path = os.path.join(OUT_DIR, f'{year}.json')
        with open(path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=1)
        print(f'{year}: {len(days):3d} days  {totals(days)}')


if __name__ == '__main__':
    main()
